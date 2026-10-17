"""Local API for StudySpace: document extraction + OpenRouter open-source LLM study-kit generation."""
from __future__ import annotations

import io
import json
import os
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from flask import Flask, jsonify, request, send_from_directory
from werkzeug.utils import secure_filename

BASE_DIR = Path(__file__).resolve().parent
HTML_FILE = BASE_DIR / "studyspace_dashboard.html"


def _load_env_file() -> None:
    """Load key-value pairs from .env if present without hard requiring dotenv."""
    candidates = [
        BASE_DIR / ".env",
        BASE_DIR.parent / ".env",
    ]
    for env_path in candidates:
        if env_path.is_file():
            try:
                with open(env_path, "r", encoding="utf-8") as f:
                    for line in f:
                        line = line.strip()
                        if line and not line.startswith("#") and "=" in line:
                            k, v = line.split("=", 1)
                            k, v = k.strip(), v.strip().strip('"\'')
                            if k not in os.environ:
                                os.environ[k] = v
            except Exception:
                pass


_load_env_file()
try:
    from dotenv import load_dotenv

    load_dotenv(BASE_DIR / ".env")
    load_dotenv(BASE_DIR.parent / ".env")
except ImportError:
    pass

OPENROUTER_API_URL = os.getenv("OPENROUTER_API_URL", "https://openrouter.ai/api/v1/chat/completions")
OPENROUTER_MODEL = os.getenv("OPENROUTER_MODEL", "nvidia/nemotron-3.5-lightning:free")
OPENROUTER_API_KEY = os.getenv("OPENROUTER_API_KEY", "")
MAX_UPLOAD_BYTES = 12 * 1024 * 1024
MAX_SOURCE_CHARS = 24_000
ALLOWED_EXTENSIONS = {"txt", "md", "csv", "pdf", "docx", "pptx"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES


def extract_text(filename: str, raw: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in {f".{x}" for x in ALLOWED_EXTENSIONS}:
        raise ValueError("Supported formats are TXT, MD, CSV, PDF, DOCX and PPTX.")
    if ext in {".txt", ".md", ".csv"}:
        text = raw.decode("utf-8-sig", errors="replace")
    elif ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(raw))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(raw))
        chunks = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(chunks)
    else:  # .pptx
        from pptx import Presentation

        deck = Presentation(io.BytesIO(raw))
        slides = []
        for number, slide in enumerate(deck.slides, start=1):
            parts = [
                shape.text
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            if parts:
                slides.append(f"Slide {number}\n" + "\n".join(parts))
        text = "\n\n".join(slides)
    text = text.strip()
    if not text:
        raise ValueError(
            "No selectable text was found. Scanned/image-only PDFs need OCR, which this version does not include."
        )
    return text[:MAX_SOURCE_CHARS]


def extract_json_from_text(text: str) -> dict:
    """Extract a JSON object from raw LLM output, handling markdown blocks or conversational preamble."""
    text = text.strip()
    if "```" in text:
        lines = text.split("\n")
        cleaned_lines = []
        inside_code = False
        for line in lines:
            trimmed = line.strip()
            if trimmed.startswith("```"):
                inside_code = not inside_code
                continue
            if inside_code:
                cleaned_lines.append(line)
        if cleaned_lines:
            candidate = "\n".join(cleaned_lines).strip()
            try:
                return json.loads(candidate)
            except json.JSONDecodeError:
                pass
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end != -1 and end > start:
        candidate = text[start : end + 1]
        try:
            return json.loads(candidate)
        except json.JSONDecodeError:
            pass
    return json.loads(text)


def call_openrouter(
    prompt: str, api_key: str | None = None, model: str | None = None
) -> dict:
    key = (api_key or os.getenv("OPENROUTER_API_KEY", "")).strip()
    if not key:
        raise RuntimeError(
            "OpenRouter API key is missing. Set OPENROUTER_API_KEY in your .env or configure it in the dashboard settings."
        )
    target_model = (model or os.getenv("OPENROUTER_MODEL", OPENROUTER_MODEL)).strip()

    payload = {
        "model": target_model,
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are StudySpace, a careful study assistant. "
                    "You create high-quality study materials strictly formatted as a valid JSON object."
                ),
            },
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.2,
        "response_format": {"type": "json_object"},
    }
    headers = {
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://127.0.0.1:5000",
        "X-Title": "StudySpace",
    }

    def _execute(body_dict: dict) -> dict:
        req = Request(
            OPENROUTER_API_URL,
            data=json.dumps(body_dict).encode("utf-8"),
            headers=headers,
            method="POST",
        )
        try:
            with urlopen(req, timeout=120) as response:
                return json.loads(response.read().decode("utf-8"))
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")[:600]
            try:
                err_data = json.loads(detail)
                msg = err_data.get("error", {}).get("message", detail)
            except Exception:
                msg = detail
            raise RuntimeError(f"OpenRouter returned HTTP {exc.code}: {msg}") from exc
        except URLError as exc:
            raise RuntimeError(
                f"Cannot connect to OpenRouter: {exc.reason}. Please check your internet connection."
            ) from exc
        except TimeoutError as exc:
            raise RuntimeError(
                "OpenRouter request timed out. Try a shorter document or a faster model."
            ) from exc

    try:
        result = _execute(payload)
    except RuntimeError as exc:
        if "response_format" in str(exc) and "response_format" in payload:
            payload.pop("response_format", None)
            result = _execute(payload)
        else:
            raise

    choices = result.get("choices", [])
    if not choices:
        raise RuntimeError("OpenRouter returned an empty response.")
    content = choices[0].get("message", {}).get("content", "")
    if not content:
        raise RuntimeError("OpenRouter model returned an empty message.")

    try:
        return extract_json_from_text(content)
    except json.JSONDecodeError as exc:
        raise RuntimeError(
            f"The model output could not be parsed as valid JSON: {exc}. Try regenerating."
        ) from exc


@app.get("/")
def index():
    return send_from_directory(BASE_DIR, HTML_FILE.name)


@app.get("/api/health")
def health():
    key = os.getenv("OPENROUTER_API_KEY", "").strip()
    model = os.getenv("OPENROUTER_MODEL", OPENROUTER_MODEL).strip()
    return jsonify(
        ok=True,
        provider="openrouter",
        model=model,
        has_api_key=bool(key),
    )


@app.post("/api/upload")
def upload():
    uploaded = request.files.get("file")
    if uploaded is None or not uploaded.filename:
        return jsonify(error="Choose a file to upload."), 400
    filename = secure_filename(uploaded.filename)
    try:
        text = extract_text(filename, uploaded.read())
        return jsonify(filename=filename, text=text, characters=len(text))
    except (ValueError, ImportError) as exc:
        return jsonify(error=str(exc)), 400
    except Exception as exc:
        app.logger.exception("Document extraction failed")
        return jsonify(error=f"Could not extract this document: {exc}"), 500


@app.post("/api/generate-study-kit")
def generate_study_kit():
    body = request.get_json(silent=True) or {}
    source = str(body.get("text", "")).strip()[:MAX_SOURCE_CHARS]
    filename = str(body.get("filename", "study material"))[:160]
    client_api_key = (
        str(body.get("apiKey", "")).strip()
        or request.headers.get("X-OpenRouter-Key", "").strip()
    )
    client_model = str(body.get("model", "")).strip()

    if len(source) < 40:
        return jsonify(
            error="Upload a document with at least 40 characters of readable text."
        ), 400

    prompt = f"""You are StudySpace, a careful study assistant. Create a useful study kit ONLY from the source text below. Do not invent facts. If a detail is unclear, avoid testing it.
Return ONLY a valid JSON object with exactly these keys:
- "notes": array of 5 to 8 objects, each with "title" and "content" strings.
- "quiz": array of 5 multiple-choice questions. Each has "question" string, "options" array of exactly 4 strings, "answer" integer index 0-3, and "explanation" string.
- "flashcards": array of 6 to 10 objects, each with "front" and "back" strings.
Make the questions test different key concepts. Ensure every answer index is correct. Use concise, clear language. Do not include markdown fences.

SOURCE FILE: {filename}
SOURCE TEXT:
{source}"""

    try:
        kit = call_openrouter(
            prompt, api_key=client_api_key or None, model=client_model or None
        )
        for key in ("notes", "quiz", "flashcards"):
            if not isinstance(kit.get(key), list) or not kit[key]:
                raise RuntimeError(f"Model output is missing a valid '{key}' list.")
        effective_model = client_model or os.getenv("OPENROUTER_MODEL", OPENROUTER_MODEL)
        return jsonify(study_kit=kit, model=effective_model, provider="openrouter")
    except RuntimeError as exc:
        return jsonify(error=str(exc)), 502
    except Exception as exc:
        app.logger.exception("Study-kit generation failed")
        return jsonify(error=f"Study-kit generation failed: {exc}"), 500


@app.errorhandler(413)
def too_large(_error):
    return jsonify(error="File is too large. The maximum upload size is 12 MB."), 413


if __name__ == "__main__":
    configured_model = os.getenv("OPENROUTER_MODEL", OPENROUTER_MODEL)
    has_key = bool(os.getenv("OPENROUTER_API_KEY", "").strip())
    print(f"StudySpace is starting with OpenRouter LLM.")
    print(f"Model: {configured_model}")
    print(f"API Key configured via environment: {'Yes' if has_key else 'No (can be configured in UI or .env)'}")
    app.run(host="127.0.0.1", port=int(os.getenv("PORT", "5000")), debug=False)
